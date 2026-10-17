"""Carga el Excel de una sede (y opcionalmente el de OC) a Supabase.

Uso:
  python -I tools/cargar_excel.py --sede SM --archivo "ruta/Backend - Tada SM.xlsx" [--catalogos] [--ejecutar] [--reset]
  python -I tools/cargar_excel.py --oc "ruta/ORDEN DE COMPRA.xlsx" [--ejecutar]

Sin --ejecutar solo valida y muestra conteos. Lee SUPABASE_URL y SUPABASE_SECRET_KEY de .env.
"""
import argparse, datetime as dt, json, os, re, sys, urllib.request, urllib.error, urllib.parse
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AVISOS = []


def aviso(msg):
    AVISOS.append(msg)


def leer_env():
    env = {}
    with open(os.path.join(RAIZ, '.env'), encoding='utf-8') as f:
        for linea in f:
            if '=' in linea and not linea.startswith('#'):
                k, v = linea.strip().split('=', 1)
                env[k] = v
    return env


class Api:
    def __init__(self, url, key):
        self.url, self.key = url.rstrip('/') + '/rest/v1/', key

    def _req(self, metodo, ruta, cuerpo=None, extra=None):
        h = {'apikey': self.key, 'Content-Type': 'application/json'}
        h.update(extra or {})
        data = json.dumps(cuerpo, ensure_ascii=False, default=str).encode('utf-8') if cuerpo is not None else None
        req = urllib.request.Request(self.url + ruta, data=data, method=metodo, headers=h)
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return r.status, r.read().decode('utf-8'), r.headers
        except urllib.error.HTTPError as e:
            raise RuntimeError(f'{metodo} {ruta} -> {e.code}: {e.read().decode("utf-8")[:400]}')

    def insertar(self, tabla, filas, lote=1000, on_conflict=None, ignorar=False):
        ruta = tabla + (f'?on_conflict={on_conflict}' if on_conflict else '')
        pref = 'return=minimal' + (',resolution=' + ('ignore' if ignorar else 'merge') + '-duplicates' if on_conflict else '')
        for i in range(0, len(filas), lote):
            self._req('POST', ruta, filas[i:i + lote], {'Prefer': pref})

    def contar(self, tabla, filtro=''):
        _, _, h = self._req('GET', f'{tabla}?select=*&limit=1{filtro}', extra={'Prefer': 'count=exact'})
        return int(h['Content-Range'].split('/')[-1])

    def obtener(self, ruta):
        return json.loads(self._req('GET', ruta)[1])

    def borrar(self, tabla, filtro):
        self._req('DELETE', f'{tabla}?{filtro}')


# ---------- normalizadores ----------
def vacio(v):
    return v is None or (isinstance(v, str) and v.strip() in ('', 'None'))


def txt(v):
    if vacio(v):
        return None
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def entero(v):
    if vacio(v):
        return None
    try:
        return int(round(float(str(v).replace(',', '.'))))
    except ValueError:
        return None


def decimal(v):
    if vacio(v):
        return None
    try:
        return round(float(str(v).replace(',', '.')), 2)
    except ValueError:
        return None


def fecha(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, str):
        s = v.strip()
        m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', s)
        if m:
            return f'{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}'
        m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
        if m:
            return f'{m[1]}-{m[2]}-{m[3]}'
    return None


def hora(v):
    if isinstance(v, dt.datetime):
        return v.time().strftime('%H:%M:%S')
    if isinstance(v, dt.time):
        return v.strftime('%H:%M:%S')
    if isinstance(v, (int, float)) and 0 <= v < 1:
        s = round(v * 86400)
        return f'{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}'
    if isinstance(v, str):
        m = re.search(r'(\d{1,2}):(\d{2})(?::(\d{2}))?', v)
        if m:
            return f'{int(m[1]):02d}:{m[2]}:{m[3] or "00"}'
    return None


def hoja_filas(wb, nombre, col_min=1, col_max=None):
    ws = wb[nombre]
    filas = ws.iter_rows(min_row=2, min_col=col_min, max_col=col_max or ws.max_column, values_only=True)
    return [f for f in filas if any(not vacio(c) for c in f)]


def ultimo_por(filas, clave):
    d = {}
    for f in filas:
        d[f[clave]] = f
    return list(d.values())


# ---------- catálogos globales ----------
def catalogos(wb, sede_id):
    maestro = {}
    for r in hoja_filas(wb, 'MAESTRO', 1, 4):
        sku = txt(r[0])
        if sku:
            maestro[sku] = {'sku': sku, 'descripcion': txt(r[1]) or '', 'cod_sap': txt(r[2]), 'precio': decimal(r[3])}

    mconteo = []
    for r in hoja_filas(wb, 'MCONTEO', 1, 5):
        ean, sku = txt(r[0]), txt(r[1])
        if ean and sku:
            mconteo.append({'sede_id': sede_id, 'ean': ean, 'sku': sku, 'sap': txt(r[2]), 'producto': txt(r[3]),
                            'multiplicador': decimal(r[4]) or 1})

    mbees = {}
    for r in hoja_filas(wb, 'MBEES', 1, 3):
        desc = txt(r[0])
        if desc:
            m = re.search(r'\d+', str(r[2])) if not vacio(r[2]) else None
            mbees[desc] = {'sede_id': sede_id, 'descripcion': desc, 'sku': txt(r[1]), 'factor': int(m[0]) if m else 1}

    env_prod, env_sku = {}, {}
    for r in hoja_filas(wb, 'MENVASES', 1, 8):
        desc, tipo = txt(r[0]), txt(r[1])
        if desc and tipo:
            if tipo not in ('Retornable', 'Venta'):
                aviso(f'MENVASES: tipo inesperado {tipo!r} en {desc!r}, omitido')
            else:
                env_prod[desc] = {'sede_id': sede_id, 'descripcion': desc, 'tipo': tipo}
        sku = txt(r[3])
        if sku:
            env_sku[sku] = {'sede_id': sede_id, 'sku': sku, 'descripcion': txt(r[4]), 'tamano': entero(r[5]),
                            'llena': txt(r[6]), 'vacio_envases': txt(r[7])}
    return {'maestro_sku': list(maestro.values()), 'mconteo': mconteo, 'mbees': list(mbees.values()),
            'envases_producto': list(env_prod.values()), 'envases_sku': list(env_sku.values())}


# ---------- datos de una sede ----------
def datos_sede(wb_valores, wb_formulas, sede_id):
    out = {}

    sp = {}
    for r in hoja_filas(wb_valores, 'STOCK', 1, 14):
        sku = txt(r[0])
        if sku:
            sp[sku] = {'sede_id': sede_id, 'sku': sku, 'descripcion': txt(r[1]) or '', 'marca': txt(r[2]),
                       'categoria': txt(r[3]), 'packs': entero(r[12]), 'top': txt(r[11])}
    out['sede_productos'] = list(sp.values())

    # ESTADO escrito a mano (no fórmula) en SALIDAS, p.ej. 'Ajuste-Inv-Jun26'
    ws_f = wb_formulas['SALIDAS']
    override = {}
    for i, row in enumerate(ws_f.iter_rows(min_row=2, min_col=11, max_col=11, values_only=True), start=2):
        v = row[0]
        if isinstance(v, str) and v.strip() and not v.startswith('='):
            override[i] = v.strip()

    ws = wb_valores['SALIDAS']
    salidas = []
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=13, values_only=True), start=2):
        if all(vacio(c) for c in r):
            continue
        f, sku = fecha(r[0]), txt(r[2])
        if not f or not sku:
            aviso(f'SALIDAS fila {i}: sin fecha o sku, omitida')
            continue
        rev = txt(r[11])
        if rev and rev != 'Entregado':
            aviso(f'SALIDAS fila {i}: revision {rev!r} -> ignorada')
            rev = None
        salidas.append({'sede_id': sede_id, 'fecha': f, 'hora': hora(r[1]), 'sku': sku, 'nro_pedido': txt(r[3]),
                        'descripcion': txt(r[4]), 'pack': entero(r[5]), 'unidad': entero(r[6]),
                        'motorizado': txt(r[8]), 'revision': rev, 'estado_override': override.get(i),
                        'origen': 'migracion'})
    out['salidas'] = salidas

    canc = []
    for i, r in enumerate(wb_valores['PD CANCELADOS'].iter_rows(min_row=2, max_col=13, values_only=True), start=2):
        if all(vacio(c) for c in r):
            continue
        nro = txt(r[3])
        if not nro:
            aviso(f'PD CANCELADOS fila {i}: sin nro de pedido, omitida')
            continue
        f = fecha(r[0])
        canc.append({'sede_id': sede_id, 'fecha': f, 'hora': hora(r[1]), 'sku': txt(r[2]), 'nro_pedido': nro,
                     'descripcion': txt(r[4]), 'pack': entero(r[5]), 'unidad': entero(r[6]), 'marca': txt(r[7]),
                     'motorizado': txt(r[8]), 'envases': txt(r[9]), 'estado': txt(r[10]), 'revision': txt(r[11]),
                     'motivo': txt(r[12]), 'cancelado_at': (f + 'T12:00:00-05:00') if f else None})
    out['cancelados'] = canc

    ent, ult_fecha = [], None
    for i, r in enumerate(wb_valores['ENTRADAS'].iter_rows(min_row=2, max_col=10, values_only=True), start=2):
        if all(vacio(c) for c in r):
            continue
        f, sku, und = fecha(r[1]), txt(r[2]), entero(r[4])
        if not f and sku and und is not None and ult_fecha:
            aviso(f'ENTRADAS fila {i}: sin fecha, se usa la de la fila anterior ({ult_fecha}) para {sku} x{und}')
            f = ult_fecha
        if not f or not sku or und is None:
            aviso(f'ENTRADAS fila {i}: falta fecha/sku/unidades ({r[1]!r}, {r[2]!r}, {r[4]!r}), omitida')
            continue
        ult_fecha = f
        fv = r[5]
        ent.append({'sede_id': sede_id, 'factura': txt(r[0]), 'fecha': f, 'sku': sku, 'descripcion': txt(r[3]),
                    'unidades': und, 'f_vencimiento': fecha(fv) or txt(fv), 'proveedor': txt(r[6]), 'oc': txt(r[7]),
                    'dynamics': txt(r[8]), 'responsable': txt(r[9])})
    out['entradas'] = ent

    mt = []
    for i, r in enumerate(wb_valores['M-T'].iter_rows(min_row=2, max_col=11, values_only=True), start=2):
        if all(vacio(c) for c in r):
            continue
        f, sku, cant, tipo = fecha(r[0]), txt(r[1]), entero(r[3]), txt(r[4])
        if not f or not sku or cant is None:
            aviso(f'M-T fila {i}: falta fecha/sku/cantidad, omitida')
            continue
        if not tipo:
            aviso(f'M-T fila {i}: sin tipo, se carga como SIN TIPO ({sku} x{cant})')
            tipo = 'SIN TIPO'
        mt.append({'sede_id': sede_id, 'fecha': f, 'sku': sku, 'descripcion': txt(r[2]), 'cantidad': cant,
                   'tipo': tipo, 'motivo': txt(r[5]), 'observaciones': txt(r[6]), 'nro_factura': txt(r[8]),
                   'fecha_retiro': fecha(r[9]), 'costo': decimal(r[10])})
    out['mt'] = mt

    cont, uids = [], set()
    for i, r in enumerate(wb_valores['CONTEO'].iter_rows(min_row=2, max_col=11, values_only=True), start=2):
        if all(vacio(c) for c in r):
            continue
        f, sku = fecha(r[0]), txt(r[3])
        if not f or not sku:
            aviso(f'CONTEO fila {i}: falta fecha o sku, omitida')
            continue
        uid = txt(r[10])
        if uid in uids:
            aviso(f'CONTEO fila {i}: uid repetido {uid!r}, se guarda sin uid')
            uid = None
        if uid:
            uids.add(uid)
        cont.append({'sede_id': sede_id, 'uid': uid, 'fecha': f, 'hora': hora(r[1]), 'ean': txt(r[2]), 'sku': sku,
                     'producto': txt(r[4]), 'packs': entero(r[5]) or 0, 'und_packs': entero(r[6]) or 0,
                     'unds': entero(r[7]) or 0, 'total': entero(r[8]) or 0, 'ubicacion': txt(r[9])})
    out['conteos'] = cont
    return out


def ordenes_compra(wb):
    filas = []
    for r in hoja_filas(wb, 'ORDEN DE COMPRA', 1, 16):
        oc = txt(r[0])
        if not oc:
            continue
        est = txt(r[15]) or 'Pendiente'
        if est not in ('Pendiente', 'Recibido', 'Cancelada'):
            aviso(f'OC {oc}: estado {est!r} -> Pendiente')
            est = 'Pendiente'
        filas.append({'oc': oc, 'cuenta_proveedor': txt(r[1]), 'cuenta_facturacion': txt(r[2]), 'proveedor': txt(r[3]),
                      'tipo_compra': txt(r[4]), 'estado_aprobacion': txt(r[5]), 'estado_orden': txt(r[6]),
                      'divisa': txt(r[7]), 'fecha_recepcion': fecha(r[8]), 'modo_entrega': txt(r[9]),
                      'condiciones_entrega': txt(r[10]), 'acuerdo_compra': txt(r[11]), 'entrega_directa': txt(r[12]),
                      'subcontrato_proyecto': txt(r[13]), 'nombre_entrega': txt(r[14]), 'estado_recepcion': est})
    return filas


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--sede')
    ap.add_argument('--archivo')
    ap.add_argument('--oc')
    ap.add_argument('--catalogos', action='store_true')
    ap.add_argument('--ejecutar', action='store_true')
    ap.add_argument('--reset', action='store_true')
    a = ap.parse_args()

    env = leer_env()
    api = Api(env['SUPABASE_URL'], env['SUPABASE_SECRET_KEY'])
    plan = {}

    if a.oc:
        plan['ordenes_compra'] = ordenes_compra(openpyxl.load_workbook(a.oc, data_only=True))

    sede_id = None
    if a.sede:
        sede = api.obtener(f'sedes?codigo=eq.{urllib.parse.quote(a.sede)}&select=id')
        if not sede:
            sys.exit(f'Sede {a.sede} no existe en la tabla sedes')
        sede_id = sede[0]['id']
        wb_v = openpyxl.load_workbook(a.archivo, data_only=True)
        wb_f = openpyxl.load_workbook(a.archivo, data_only=False)
        if a.catalogos:
            plan.update(catalogos(wb_v, sede_id))
        plan.update(datos_sede(wb_v, wb_f, sede_id))

    print('--- Plan ---')
    for t, filas in plan.items():
        print(f'{t:18s} {len(filas):>7d} filas')
    print(f'--- Avisos ({len(AVISOS)}) ---')
    for m in AVISOS[:40]:
        print(' ', m)
    if len(AVISOS) > 40:
        print(f'  ... y {len(AVISOS) - 40} más')

    if not a.ejecutar:
        print('\nSimulación: no se escribió nada. Agrega --ejecutar para cargar.')
        return

    por_sede = ['sede_productos', 'salidas', 'cancelados', 'entradas', 'mt', 'conteos']
    cat_sede = ['mbees', 'mconteo', 'envases_producto', 'envases_sku']
    if sede_id:
        for t in por_sede:
            n = api.contar(t, f'&sede_id=eq.{sede_id}')
            if n and not a.reset:
                sys.exit(f'{t} ya tiene {n} filas para la sede; usa --reset para reemplazarlas')
        if a.reset:
            for t in reversed(por_sede):
                api.borrar(t, f'sede_id=eq.{sede_id}')

    orden = ['maestro_sku', 'mconteo', 'mbees', 'envases_producto', 'envases_sku', 'sede_productos',
             'salidas', 'cancelados', 'entradas', 'mt', 'conteos', 'ordenes_compra']
    claves = {'maestro_sku': 'sku', 'mbees': 'sede_id,descripcion', 'envases_producto': 'sede_id,descripcion', 'envases_sku': 'sede_id,sku',
              'ordenes_compra': 'oc', 'sede_productos': 'sede_id,sku'}
    for t in orden:
        if t in plan and plan[t]:
            api.insertar(t, plan[t], on_conflict=claves.get(t), ignorar=(t == 'maestro_sku'))
            print(f'cargado {t}: {len(plan[t])}')
    print('Listo.')


if __name__ == '__main__':
    main()
